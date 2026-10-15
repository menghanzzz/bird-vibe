#!/usr/bin/env python3
"""测试集构建脚本：从 training_data/val 分层抽样构建评测集（含泄漏检查硬门槛）。

设计要点（对应评测方案 Phase 1）：
- 核心层：每类从 val/ 确定性抽 PER_CLASS 张（固定种子，可复现），金标 = 物种目录名
- 野外层：复用 eval_photos/ 现有 12 张实拍（含已有人工金标）
- 【硬门槛】数据泄漏检查（不通过则整体拒绝构建）：
  1. manifest observation_id 级：选中图的 observation_id 不得出现在同物种 train/ 中
  2. 文件名级兜底：val 文件名（observation_id）不得与 train/ 文件名重叠
  3. 全局 sha256 去重：跨物种重复文件视为数据缺陷，硬失败
- 产出：eval_set/photos/ + eval_set/testset.xlsx（列兼容 batch_eval.py）+ build_manifest.json（审计）

用法：
  backend/venv/bin/python backend/scripts/build_eval_set.py              # 默认每类 7 张
  backend/venv/bin/python backend/scripts/build_eval_set.py --per-class 6
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
import shutil
from datetime import datetime, timezone
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
TRAIN_DATA = BACKEND / "training_data"
MANIFEST = TRAIN_DATA / "manifest.json"
OUT_DIR = BACKEND / "eval_set"
PHOTOS_DIR = OUT_DIR / "photos"
TESTSET_XLSX = OUT_DIR / "testset.xlsx"
BUILD_MANIFEST = OUT_DIR / "build_manifest.json"
LEGACY_XLSX = BACKEND / "eval_photos.xlsx"
LEGACY_PHOTOS = BACKEND / "eval_photos"
KB_DIR = BACKEND / "knowledge_base"

SEED = 42

# 与 batch_eval.py 的列约定保持一致（它按列名读取，多余列忽略）
COLUMNS = [
    "照片", "实际鸟种",
    "带时空 top3", "不带 top3", "top1 命中", "top3 命中",
    "bioclip top3", "bioclip 命中", "bioclip_tuned top3", "bioclip_tuned 命中",
    "层", "类别", "集", "photo_id",
]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def kb_categories() -> dict[str, str]:
    """从知识库 frontmatter 读类别（非 KB 物种返回空串）。"""
    import yaml
    cats = {}
    for p in KB_DIR.glob("*.md"):
        try:
            fm = yaml.safe_load(p.read_text(encoding="utf-8").split("---")[1])
            cats[p.stem] = (fm or {}).get("category", "")
        except Exception:
            cats[p.stem] = ""
    return cats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--per-class", type=int, default=7)
    ap.add_argument("--seed", type=int, default=SEED)
    args = ap.parse_args()

    if not MANIFEST.exists():
        raise SystemExit(f"找不到 {MANIFEST}")
    species_manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))["species"]
    cats = kb_categories()
    kb = {p.stem for p in KB_DIR.glob("*.md")}
    species_dirs = sorted([d for d in TRAIN_DATA.iterdir() if d.is_dir()])
    print(f"[构建] 物种目录 {len(species_dirs)} 个，每类抽 {args.per_class} 张，seed={args.seed}")

    # ── 第 1 步：抽样 + 泄漏检查 + 去重（单次遍历，产出选中清单）──
    selections = []          # (src_path, species, obs_id)
    violations: list[str] = []
    hash_seen: dict[str, str] = {}
    audit: dict[str, dict] = {}

    for sp in species_dirs:
        val_dir = sp / "val"
        val_imgs = sorted(val_dir.glob("*.jpg")) if val_dir.is_dir() else []
        if len(val_imgs) < 5:
            print(f"[构建] ⚠️ {sp.name} val 仅 {len(val_imgs)} 张（<5），全部取用")
        rng = random.Random(f"{args.seed}:{sp.name}")
        pool = val_imgs[:]
        rng.shuffle(pool)
        selected = pool[:args.per_class]

        # 泄漏检查
        images = species_manifest.get(sp.name, {}).get("images", [])
        obs_by_file = {img["file"]: img.get("observation_id") for img in images}
        train_obs = {img.get("observation_id") for img in images
                     if img.get("file", "").startswith(f"{sp.name}/train/")}
        train_dir = sp / "train"
        train_fnames = {p.name for p in train_dir.glob("*")} if train_dir.is_dir() else set()

        kept = 0
        for p in selected:
            rel = f"{sp.name}/val/{p.name}"
            obs_id = obs_by_file.get(rel)
            if obs_id is None:
                violations.append(f"{sp.name}/{p.name}: manifest 中找不到对应条目")
                continue
            if obs_id in train_obs:
                violations.append(f"{sp.name}/{p.name}: observation_id {obs_id} 同时出现在 train/")
            if p.name in train_fnames:
                violations.append(f"{sp.name}/{p.name}: 文件名与 train/ 重叠")
            digest = sha256_file(p)
            if digest in hash_seen:
                violations.append(f"{sp.name}/{p.name}: 与 {hash_seen[digest]} 完全重复（sha256）")
                continue
            hash_seen[digest] = f"{sp.name}/{p.name}"
            selections.append((p, sp.name, obs_id))
            kept += 1
        audit[sp.name] = {"val_total": len(val_imgs), "selected": kept}

    # ── 硬门槛：任何泄漏/数据缺陷 → 拒绝构建 ──
    if violations:
        OUT_DIR.mkdir(exist_ok=True)
        (OUT_DIR / "leakage_violations.txt").write_text(
            "\n".join(violations), encoding="utf-8")
        raise SystemExit(
            f"⛔ 泄漏检查未通过：{len(violations)} 处违规"
            f"（明细见 eval_set/leakage_violations.txt），已拒绝构建测试集。")

    # ── 第 2 步：写产物（复制照片 + Excel + 审计 manifest）──
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    PHOTOS_DIR.mkdir(parents=True)

    rows = []
    for i, (src, species, obs_id) in enumerate(selections, 1):
        fname = f"c{i:03d}_{species}.jpg"
        shutil.copy2(src, PHOTOS_DIR / fname)
        rows.append((fname, species, "core", cats.get(species, "非KB"), "val", str(obs_id)))

    from openpyxl import Workbook, load_workbook
    legacy_wb = load_workbook(LEGACY_XLSX)
    legacy_ws = legacy_wb.active
    header = {c.value: i + 1 for i, c in enumerate(legacy_ws[1])}
    wild_n = 0
    for r in range(2, legacy_ws.max_row + 1):
        photo = legacy_ws.cell(r, header["照片"]).value
        actual = legacy_ws.cell(r, header["实际鸟种"]).value
        if not photo or not actual:
            continue
        src = LEGACY_PHOTOS / str(photo).strip()
        if not src.exists():
            continue
        wild_n += 1
        fname = f"w{wild_n:02d}_{Path(str(photo).strip()).stem}.jpg"
        shutil.copy2(src, PHOTOS_DIR / fname)
        label = str(actual).strip()
        rows.append((fname, label, "wild", cats.get(label, "非KB"), "legacy", ""))

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.append(COLUMNS)
    for r in rows:
        out_ws.append(list(r) + [""] * (len(COLUMNS) - len(r)))
    out_wb.save(TESTSET_XLSX)

    non_kb = sorted({r[1] for r in rows if r[2] == "core" and r[1] not in kb})
    BUILD_MANIFEST.write_text(json.dumps({
        "built_at": datetime.now(timezone.utc).isoformat(),
        "seed": args.seed,
        "per_class": args.per_class,
        "leakage_gate": {"observation_id_level": True, "filename_level": True,
                          "global_sha256_dedup": True, "violations": 0},
        "species": audit,
        "totals": {"core": sum(1 for r in rows if r[2] == "core"),
                    "wild": wild_n, "total": len(rows),
                    "non_kb_classes": non_kb},
        "source_manifest_sha256": sha256_file(MANIFEST),
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    short = [f"{s}({a['selected']})" for s, a in audit.items() if a["selected"] < args.per_class]
    print(f"[构建] ✅ 完成：核心层 {sum(1 for r in rows if r[2]=='core')} + 野外层 {wild_n} "
          f"= {len(rows)} 张 -> {PHOTOS_DIR}")
    print(f"[构建] 非 KB 物种 {len(non_kb)} 个（用于测试库外行为）: {non_kb}")
    if short:
        print(f"[构建] ⚠️ 未达每类 {args.per_class} 张: {short}")
    print(f"[构建] 泄漏检查: 0 违规（observation_id / 文件名 / sha256 三道门槛全部通过）")
    print(f"[构建] 产物: {TESTSET_XLSX} | {BUILD_MANIFEST}")


if __name__ == "__main__":
    main()
